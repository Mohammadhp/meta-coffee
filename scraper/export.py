#!/usr/bin/env python
"""Export catalog.jsonl to CSV + XLSX."""
import csv
import json

import openpyxl
from openpyxl.utils import get_column_letter

IN = "scraper/catalog.jsonl"
CSV_OUT = "scraper/catalog.csv"
BEANS_CSV_OUT = "scraper/catalog-beans.csv"
XLSX_OUT = "scraper/catalog.xlsx"

COLUMNS = [
    ("roaster", "Roaster"),
    ("product_name", "Product Name"),
    ("origin", "Origin"),
    ("process", "Process"),
    ("roast_level", "Roast Level"),
    ("format", "Format"),
    ("weight_g", "Weight (g)"),
    ("price_toman", "Price (Toman)"),
    ("price_per_100g", "Price / 100g"),
    ("specialty_score", "Specialty Score"),
    ("in_stock", "In Stock?"),
    ("categories", "Categories"),
    ("product_url", "Product URL"),
    ("description", "Description"),
]


def load():
    recs = [json.loads(l) for l in open(IN, encoding="utf-8")]
    return recs


def to_row(rec):
    return [rec.get(k) for k, _ in COLUMNS]


def is_bean(rec):
    """Bean = origin-detected. Gear/accessories/chocolate/etc. have no origin.
    Not perfect (see false-positive note), but the right starting set to curate."""
    return rec.get("origin") is not None


def write_csv(recs, path, header=True):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        if header:
            w.writerow([h for _, h in COLUMNS])
        for r in recs:
            w.writerow(to_row(r))


def _add_sheet(wb, title, recs):
    ws = wb.create_sheet(title=title)
    ws.append([h for _, h in COLUMNS])
    for r in recs:
        ws.append(to_row(r))
    for c in range(1, len(COLUMNS) + 1):
        ws.cell(row=1, column=c).font = openpyxl.styles.Font(bold=True)
        ws.column_dimensions[get_column_letter(c)].width = 22
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["M"].width = 50
    ws.column_dimensions["N"].width = 60
    ws.freeze_panes = "A2"
    return ws


def write_xlsx(recs, beans):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Catalogue Items"
    ws.append([h for _, h in COLUMNS])
    for r in recs:
        ws.append(to_row(r))
    for c in range(1, len(COLUMNS) + 1):
        ws.cell(row=1, column=c).font = openpyxl.styles.Font(bold=True)
        ws.column_dimensions[get_column_letter(c)].width = 22
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["M"].width = 50
    ws.column_dimensions["N"].width = 60
    ws.freeze_panes = "A2"
    _add_sheet(wb, "Beans", beans)
    wb.save(XLSX_OUT)


def main():
    recs = load()
    beans = [r for r in recs if is_bean(r)]
    write_csv(recs, CSV_OUT)
    write_csv(beans, BEANS_CSV_OUT)
    write_xlsx(recs, beans)
    print(f"Exported {len(recs)} records ({len(beans)} beans)")
    print(f"  -> {CSV_OUT}")
    print(f"  -> {BEANS_CSV_OUT}")
    print(f"  -> {XLSX_OUT}")


if __name__ == "__main__":
    main()
